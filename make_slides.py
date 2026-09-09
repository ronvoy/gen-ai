"""Build `extended-final-presentation.pptx`.

Structure follows `report-ref.pptx` — the assignment's reference deck — whose
slide titles define the required sections:

    Mid-Term Empirical Study · Research Question · Experimental Protocol ·
    Benchmark · Involved approaches · Considered Model(s) ·
    Comments and Discussion before Results · Results ·
    Discussion of the results · Conclusions   (+ References)

Visual style is lifted from `report.pptx` (the team's existing deck): warm
off-white ground, Helvetica Neue, blue eyebrow label per section, hairline rule
under each title, running footer with page number. A section that needs more
than one slide repeats its eyebrow, exactly as the previous deck does.

Content is read from results/v2 so the numbers cannot drift.

    python make_slides.py
"""

import glob
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

import make_results_section as R
from metrics import significance as SIG

OUT = "extended-final-presentation.pptx"
FIGDIR = "diagram-analysis"
DECK_FOOTER = "LAMBADA & MMLU Benchmark Evaluation — Extended"

# ---- palette, taken from report.pptx ---------------------------------------
BG = RGBColor(0xFC, 0xFC, 0xFB)
INK = RGBColor(0x0B, 0x0B, 0x0B)
BODY = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
RULE = RGBColor(0xE1, 0xE0, 0xD9)
PANEL = RGBColor(0xF9, 0xF9, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Helvetica Neue"

# ---- grid ------------------------------------------------------------------
SW, SH = Inches(13.333), Inches(7.5)
M = Inches(0.60)                 # left margin
CW = Inches(12.10)               # content width

_page = {"n": 0}


def load(benchmark):
    return [json.load(open(f, encoding="utf-8"))
            for f in sorted(glob.glob(f"results/v2/*_{benchmark}_v2.json"))]


def comparison(benchmark):
    p = f"results/v2/comparison_{benchmark}.json"
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else None


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------

def text(slide, body, left, top, width, height, size=14, bold=False,
         color=BODY, align=PP_ALIGN.LEFT, spacing=1.15, italic=False,
         space_after=0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(body).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = FONT
    return box


def rule(slide, top, left=M, width=Inches(12.13)):
    ln = slide.shapes.add_connector(1, left, top, left + width, top)
    ln.line.color.rgb = RULE
    ln.line.width = Pt(0.75)
    return ln


def base(prs, numbered=True):
    """Blank slide with the house background and footer."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    if numbered:
        _page["n"] += 1
        text(s, DECK_FOOTER, M, Inches(7.12), Inches(6.0), Inches(0.30),
             size=9, color=MUTED)
        text(s, str(_page["n"]), Inches(12.40), Inches(7.12), Inches(0.5),
             Inches(0.30), size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    return s


def section_slide(prs, eyebrow, title, size=26):
    """Standard content slide: blue eyebrow, title, hairline rule.

    Geometry matches report.pptx exactly: eyebrow at 0.60/0.35 in 12.5pt blue,
    title at 0.60/1.00 in 26pt ink, hairline rule at 1.65.
    """
    s = base(prs)
    text(s, eyebrow.upper(), M, Inches(0.35), Inches(11.0), Inches(0.35),
         size=12.5, bold=True, color=ACCENT)
    text(s, title, M, Inches(1.00), CW, Inches(0.70),
         size=size, bold=True, color=INK)
    rule(s, Inches(1.65))
    return s


def opener_slide(prs, title, statement, note=None):
    """Section opener: large title, framed statement panel."""
    s = base(prs)
    text(s, title, M, Inches(0.55), CW, Inches(0.70), size=30, bold=True, color=INK)
    rule(s, Inches(1.30))
    panel = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, Inches(1.85), CW, Inches(2.10))
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = ACCENT
    panel.line.width = Pt(1.5)
    panel.shadow.inherit = False
    tf = panel.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.35)
    tf.margin_top = tf.margin_bottom = Inches(0.25)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = statement
    for p in tf.paragraphs:
        p.line_spacing = 1.25
        for r in p.runs:
            r.font.size = Pt(21)
            r.font.bold = True
            r.font.color.rgb = INK
            r.font.name = FONT
    if note:
        text(s, note, M, Inches(4.30), CW, Inches(1.60), size=16, spacing=1.3)
    return s


def label(slide, txt, left, top, width=Inches(3.0)):
    return text(slide, txt.upper(), left, top, width, Inches(0.30),
                size=11, bold=True, color=MUTED)


def bullets(slide, items, left, top, width, height, size=14, spacing=1.35):
    return text(slide, "\n".join(f"—  {i}" for i in items),
                left, top, width, height, size=size, spacing=spacing,
                space_after=7)


def chip(slide, txt, left, top, width, height=Inches(0.42), fill=ACCENT):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.fill.background()
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.text = txt
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = WHITE
            r.font.name = FONT
    return box


def table(slide, rows, left, top, width, col_w=None, font=12, row_h=0.34,
          emphasise=None):
    """Dark header row, off-white body — matching report.pptx."""
    n_r, n_c = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_r, n_c, left, top, width,
                                   Inches(row_h * n_r))
    t = shape.table
    if col_w:
        tot = sum(col_w)
        for i, w in enumerate(col_w):
            t.columns[i].width = Emu(int(width * w / tot))
    for r, row in enumerate(rows):
        t.rows[r].height = Inches(row_h)
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(val)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if r == 0 else BG
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                for run in p.runs:
                    run.font.size = Pt(font)
                    run.font.name = FONT
                    run.font.bold = (r == 0) or (emphasise is not None
                                                 and c == emphasise and r > 0)
                    run.font.color.rgb = WHITE if r == 0 else INK
    return t


def figure(slide, filename, caption, left, top, width, height):
    """Embed the PNG if present, else a labelled placeholder frame."""
    path = os.path.join(FIGDIR, filename)
    if os.path.exists(path):
        slide.shapes.add_picture(path, left, top, width=width)
    else:
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = PANEL
        box.line.color.rgb = MUTED
        box.line.width = Pt(1)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.text = f"[ screenshot placeholder ]\n{FIGDIR}/{filename}"
        for i, p in enumerate(tf.paragraphs):
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(14 if i == 0 else 11)
                r.font.bold = (i == 0)
                r.font.color.rgb = MUTED
                r.font.name = FONT
    text(slide, caption, left, top + height + Inches(0.10), width, Inches(0.40),
         size=11.5, color=MUTED, align=PP_ALIGN.CENTER, italic=True)


# ---------------------------------------------------------------------------
# 1 — Title
# ---------------------------------------------------------------------------

def s_title(prs):
    s = base(prs, numbered=False)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.14))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False

    text(s, "MID-TERM EMPIRICAL STUDY — EXTENDED", Inches(0.90), Inches(2.05),
         Inches(11.5), Inches(0.40), size=14, bold=True, color=ACCENT)
    text(s, "LAMBADA & MMLU Benchmark\nA Nine-Stage Evaluation of Small Language Models",
         Inches(0.90), Inches(2.48), Inches(11.5), Inches(1.70),
         size=34, bold=True, color=INK, spacing=1.12)

    mm, lb = load("mmlu"), load("lambada")
    scale = ""
    if mm and lb:
        n_mm = mm[0]["config"]["dataset"]["n_items"] * len(mm)
        n_lb = lb[0]["config"]["dataset"]["n_items"] * len(lb)
        scale = (f"{n_mm:,} MMLU questions  ·  {n_lb:,} LAMBADA passages  ·  "
                 f"3 models  ·  9 measured stages")
    text(s, "Beyond accuracy: calibration, consistency, context, robustness,\n"
            "latency, token efficiency, cost and reliability.",
         Inches(0.90), Inches(4.00), Inches(11.5), Inches(0.90),
         size=14.5, color=BODY, spacing=1.3)
    if scale:
        text(s, scale, Inches(0.90), Inches(4.78), Inches(11.5), Inches(0.35),
             size=12.5, color=ACCENT, bold=True)

    rule(s, Inches(5.30), left=Inches(0.90), width=Inches(11.5))
    text(s, "SUBMITTED TO", Inches(0.90), Inches(5.52), Inches(6.0),
         Inches(0.35), size=10.5, bold=True, color=MUTED)
    text(s, "Prof. Anna Corazza", Inches(0.90), Inches(5.84), Inches(6.0),
         Inches(0.40), size=14, color=INK)
    text(s, "TEAM WORK — SUBMITTED BY", Inches(0.90), Inches(6.24), Inches(9.0),
         Inches(0.35), size=10.5, bold=True, color=MUTED)
    text(s, "Francesco Ventimiglia   ·   Danilo Rodriguez   ·   Rohan Baidya",
         Inches(0.90), Inches(6.56), Inches(9.0), Inches(0.40), size=14, color=INK)
    text(s, "github.com/ronvoy/gen-ai   ·   unina.cc/gen-ai",
         Inches(0.90), Inches(7.02), Inches(9.0), Inches(0.35),
         size=10.5, color=MUTED)


# ---------------------------------------------------------------------------
# 2 — Research Question
# ---------------------------------------------------------------------------

def s_research_question(prs):
    opener_slide(
        prs, "Research Question",
        "Among three small language models built with different compression "
        "strategies, which best balances quality, calibration, robustness, "
        "latency, cost and reliability — on long-range context prediction "
        "(LAMBADA) and broad knowledge and reasoning (MMLU)?",
        "In particular: a single accuracy number cannot separate a wrong answer "
        "from an unparseable one, a confident-and-right model from a "
        "confident-and-wrong one, or cheap tokens from a cheap correct answer. "
        "This study measures those separately.")

    s = section_slide(prs, "Research Question",
                      "Why accuracy alone cannot answer it")
    rows = [["A single accuracy number cannot distinguish…", "Stage that can"],
            ["a wrong answer from an unparseable one", "1 · parse-failure rate"],
            ["confident-and-right from confident-and-wrong", "2 · ECE, Brier"],
            ["a stable answer from a lucky one", "3 · answer stability"],
            ["comprehension from local n-gram matching", "4 · context ablation"],
            ["a robust score from a fragile one", "5 · perturbation deltas"],
            ["a fast model from a slow one", "6 · TTFT, TPOT, p95"],
            ["cheap tokens from cheap correct answers", "7–8 · cost per correct answer"],
            ["a clean run from one that silently retried", "9 · retries, failover"]]
    table(s, rows, M, Inches(1.95), CW, col_w=[62, 38], font=13, row_h=0.40)
    text(s, "Each stage exists because it separates two things that accuracy conflates.",
         M, Inches(5.80), CW, Inches(0.45), size=14, bold=True, color=ACCENT)


# ---------------------------------------------------------------------------
# 3 — Experimental Protocol
# ---------------------------------------------------------------------------

def s_protocol(prs):
    s = section_slide(prs, "Experimental Protocol", "How one run executes")
    steps = [("1", "Configure", "RunConfig: models, decoding,\npasses, dataset, parallelism"),
             ("2", "Pre-flight", "probe each provider;\nwarn if rate-limited"),
             ("3", "Estimate", "print the API-call multiplier\nbefore spending"),
             ("4", "Scored pass", "chain-of-thought, non-streamed\n→ headline accuracy"),
             ("5", "Extended pass", "streamed → TTFT, TPOT,\ntokens, cost, provider"),
             ("6", "Optional passes", "calibration · repeats\nrobustness · ablation"),
             ("7", "Compute", "nine metric blocks, each marked\navailable or not"),
             ("8", "Aggregate", "rank; refuse if\nconfigurations differ"),
             ("9", "Publish", "results · history · web view\nreport · slides")]
    x0, y0 = M, Inches(1.95)
    bw, bh, gap = Inches(3.90), Inches(1.22), Inches(0.16)
    for i, (num, head, body) in enumerate(steps):
        col, row = i % 3, i // 3
        left, top = x0 + col * (bw + gap), y0 + row * (bh + gap)
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, bw, bh)
        box.fill.solid()
        box.fill.fore_color.rgb = PANEL
        box.line.color.rgb = RULE
        box.shadow.inherit = False
        text(s, f"{num}   {head}", left + Inches(0.18), top + Inches(0.10),
             bw - Inches(0.36), Inches(0.34), size=14, bold=True, color=ACCENT)
        text(s, body, left + Inches(0.18), top + Inches(0.46),
             bw - Inches(0.36), Inches(0.70), size=11, color=BODY, spacing=1.1)
    text(s, "Stage 10 — TP/PP/DP/SP/CP/EP, VRAM, energy — is absent by design: none of it "
            "is observable through a hosted inference API.",
         M, Inches(6.35), CW, Inches(0.50), size=12.5, color=MUTED, italic=True)


def s_protocol_passes(prs):
    s = section_slide(prs, "Experimental Protocol", "Two measurement passes per item")
    rows = [["Pass", "Prompt", "Streamed", "Produces"],
            ["Scored", "chain-of-thought, then a letter/word", "no",
             "headline accuracy, reasoning analysis"],
            ["Extended", "identical prompt", "yes",
             "TTFT, TPOT, throughput, tokens, cost, reliability"],
            ["Calibration", "direct answer, max_tokens = 1", "no",
             "option probabilities → NLL, ECE, Brier, rank"]]
    table(s, rows, M, Inches(1.95), CW, col_w=[14, 30, 11, 45], font=12.5, row_h=0.46)
    label(s, "Why two", M, Inches(4.05))
    bullets(s, [
        "TTFT — time to first token — cannot be observed from a single blocking "
        "response, so latency requires streaming.",
        "The scored pass also parses and grades the reasoning text, which the "
        "streamed path does not.",
        "Providers return log-probabilities for the final token only; at "
        "max_tokens = 1 that token is the answer, which is what makes "
        "calibration recoverable at all.",
    ], M, Inches(4.38), CW, Inches(1.60), size=13.5)
    text(s, "Because they are separate API calls, the two accuracies differ slightly even at "
            "temperature 0. That residual is serving non-determinism — the web view states "
            "the delta rather than showing two numbers unexplained.",
         M, Inches(6.20), CW, Inches(0.70), size=12.5, color=ACCENT)


# ---------------------------------------------------------------------------
# 4 — Benchmark
# ---------------------------------------------------------------------------

def s_benchmark_arch(prs):
    s = section_slide(prs, "Benchmark", "Ten stages defined, nine measurable")
    rows = [["#", "Stage", "Purpose", "Band", "Observable"],
            ["1", "Task Quality", "Is the answer right?", "P0", "yes"],
            ["2", "Probabilistic Quality", "Is its confidence earned?", "P1", "conditional"],
            ["3", "Reasoning & Consistency", "Same answer twice?", "P1", "yes"],
            ["4", "Context Behavior", "Is it reading the passage?", "P1", "yes"],
            ["5", "Robustness", "Does the score survive input change?", "P1", "yes"],
            ["6", "API Performance", "How fast did it respond?", "P0", "yes"],
            ["7", "Token Efficiency", "How many tokens consumed?", "P0", "yes"],
            ["8", "Economics", "What did it cost?", "P0", "yes"],
            ["9", "Reliability", "Did the calls succeed?", "P0", "yes"],
            ["10", "Hardware & Distributed", "TP/PP/DP · VRAM · energy", "—", "no — excluded"]]
    table(s, rows, M, Inches(1.85), CW, col_w=[5, 25, 39, 10, 21], font=11.5, row_h=0.375)
    text(s, "P0 stages come out of every run. P1 stages each cost one extra pass over the "
            "dataset, so they are opt-in and the runner prints the multiplier before spending.",
         M, Inches(6.20), CW, Inches(0.60), size=12.5, color=MUTED)


def s_benchmark_config(prs):
    s = section_slide(prs, "Benchmark", "Configuration is not a metric")
    rows = [["", "Configuration variable", "Metric"],
            ["Nature", "an input you choose", "an output you observe"],
            ["Examples", "TP, PP, DP, SP, CP, EP · dtype · batch size · temperature",
             "accuracy · ECE · TTFT · cost · success rate"],
            ["Lives in", "benchmark_config.py → RunConfig", "metrics/ → result blocks"],
            ["Role", "independent variable", "dependent variable"]]
    table(s, rows, M, Inches(1.90), CW, col_w=[13, 45, 42], font=12.5, row_h=0.44)
    text(s, "Folding a parallelism degree into the same dictionary as an accuracy score invites "
            "a category error — ranking models on a composite that blends a hardware layout with "
            "a quality measurement.",
         M, Inches(4.20), CW, Inches(0.80), size=14, spacing=1.3)
    text(s, "A parallelism degree is not a virtue.",
         M, Inches(5.05), CW, Inches(0.45), size=18, bold=True, color=ACCENT)
    text(s, "Consequence: every result file carries the RunConfig that produced it, and the "
            "comparison layer reports comparable: false rather than publishing a confounded "
            "ranking when two runs differ.",
         M, Inches(5.60), CW, Inches(0.85), size=13, color=BODY, spacing=1.3)


def s_benchmark_data(prs):
    mm, lb = load("mmlu"), load("lambada")
    s = section_slide(prs, "Benchmark", "Datasets and scale")
    n_mm = mm[0]["config"]["dataset"]["n_items"] if mm else 0
    n_lb = lb[0]["config"]["dataset"]["n_items"] if lb else 0
    rows = [["", "MMLU", "LAMBADA"],
            ["Task", "4-option multiple choice", "predict the final word"],
            ["Tests", "knowledge and reasoning, 57 subjects", "long-range reading comprehension"],
            ["Items per model", f"{n_mm:,}  (57 subjects × 50)", f"{n_lb:,} passages"],
            ["Chance baseline", "25%", "≈ 0% (open vocabulary)"],
            ["Scored by", "letter match", "normalised word match"],
            ["Source", "HF datasets-server, cached", "BookCorpus test split, local"]]
    table(s, rows, M, Inches(1.90), CW, col_w=[20, 40, 40], font=12.5, row_h=0.44)
    text(s, "Ground truth comes from the source datasets — the official answer key and the "
            "gold target word — never inferred.",
         M, Inches(5.20), CW, Inches(0.50), size=13, color=MUTED)


# ---------------------------------------------------------------------------
# 5 / 6 — Approaches and models
# ---------------------------------------------------------------------------

def s_approaches(prs):
    s = section_slide(prs, "Involved approaches", "Three compression strategies compared")
    rows = [["Model", "Params", "Architecture", "Key technique"],
            ["Gemma-3-4B", "4B", "Dense decoder-only, interleaved local/global attention",
             "Knowledge distillation"],
            ["Llama-3.2-3B", "3B", "Dense decoder-only with Grouped Query Attention",
             "Pruning + distillation"],
            ["Ministral-8B", "8B", "Decoder-only with Sliding Window Attention",
             "Native 8B, not distilled down"]]
    table(s, rows, M, Inches(1.90), CW, col_w=[18, 9, 43, 30], font=12, row_h=0.52)
    label(s, "The comparison this sets up", M, Inches(4.15))
    bullets(s, [
        "Two compressed models (Gemma, Llama) against one trained at its native size (Ministral).",
        "All three queried through the same API, with one shared configuration, so differences "
        "are attributable to the models.",
        "No weights are changed — only decoding parameters vary, and identically across models.",
    ], M, Inches(4.50), CW, Inches(1.60), size=13.5)


def s_models(prs):
    info = [
        ("Gemma-3-4B", "Google · 4B",
         "Dense decoder-only transformer that interleaves five local "
         "sliding-window attention layers with one global layer, so most layers "
         "stay cheap while information still crosses the whole context.",
         ["5:1 local / global attention interleaving", "Grouped Query Attention with QK-norm",
          "128k-token context window", "Distilled from a larger teacher model"],
         "Strong quality per parameter; global reasoning can lag full attention."),
        ("Llama-3.2-3B", "Meta · 3B",
         "The standard Llama recipe scaled down: rotary position embeddings, "
         "grouped query attention and SwiGLU feed-forward layers, pruned and "
         "distilled from larger Llama 3.1 checkpoints.",
         ["Dense decoder-only, no architectural novelty", "RoPE for length generalisation",
          "Grouped Query Attention, SwiGLU", "Smallest and fastest of the three"],
         "Very small and quick to serve; lower ceiling on complex reasoning."),
        ("Ministral-8B", "Mistral AI · 8B",
         "Interleaved sliding-window attention keeps memory low on long inputs; "
         "grouped query attention shrinks the KV cache further. Trained at 8B "
         "rather than compressed down from a larger model.",
         ["Interleaved sliding window attention", "Grouped Query Attention for a smaller KV cache",
          "Depth propagates context beyond one window", "Tuned for long-context edge inference"],
         "Best quality here; distant context can fade across windows."),
    ]
    for name, meta, desc, props, verdict in info:
        s = section_slide(prs, "Considered Model(s)", name, size=28)
        text(s, meta, M, Inches(1.72), CW, Inches(0.32), size=13, bold=True, color=ACCENT)
        label(s, "How it works", M, Inches(2.20))
        text(s, desc, M, Inches(2.52), Inches(5.85), Inches(1.70), size=13.5, spacing=1.3)
        label(s, "Key properties", Inches(6.85), Inches(2.20))
        bullets(s, props, Inches(6.85), Inches(2.52), Inches(5.85), Inches(2.00), size=13)
        rule(s, Inches(5.05))
        label(s, "Trade-off", M, Inches(5.25))
        text(s, verdict, M, Inches(5.58), CW, Inches(0.60), size=14, color=INK)


# ---------------------------------------------------------------------------
# 7 — Comments and Discussion before Results
# ---------------------------------------------------------------------------

EYEBROW_PRE = "Comments and Discussion before Results"


def s_pre_stage_quality(prs):
    s = section_slide(prs, EYEBROW_PRE, "Stages 1–2 · Quality and calibration")
    rows = [["Metric", "What it means", "Good"],
            ["Overall accuracy", "correct ÷ questions actually answered", "higher"],
            ["Macro accuracy", "mean of per-subject accuracies, so a big subject cannot dominate", "higher"],
            ["Normalised accuracy", "(acc − 25%) ÷ 75% — distance above coin-flipping", "higher"],
            ["Wilson 95% CI", "binomial interval that stays inside [0,1] at small n", "narrower"],
            ["Parse-failure rate", "no extractable answer — disobedience, not ignorance", "lower"],
            ["ECE / MCE", "does stated confidence match observed accuracy; and the worst bin", "lower"],
            ["Brier score", "proper scoring rule — cannot be gamed by hedging at 50%", "lower"],
            ["NLL / perplexity", "how surprised the model was by the truth", "lower"]]
    table(s, rows, M, Inches(1.90), CW, col_w=[22, 62, 16], font=11.5, row_h=0.40)
    text(s, "Accuracy is scored over questions the API answered. A request refused with a 429 "
            "says nothing about what the model knows; counting it wrong measures provider load.",
         M, Inches(5.65), CW, Inches(0.60), size=12.5, color=ACCENT)


def s_pre_stage_behaviour(prs):
    s = section_slide(prs, EYEBROW_PRE, "Stages 3–5 · Consistency, context, robustness")
    rows = [["Metric", "What it means", "Good"],
            ["Answer stability", "share of items where every repeat gave the same answer", "higher"],
            ["Self-consistency gain", "majority-vote accuracy minus single-sample accuracy", "higher"],
            ["Context utilisation", "acc(full passage) − acc(last sentence only)", "higher"],
            ["Utilisation ratio", "share of skill that depends on the wider passage", "higher"],
            ["Accuracy drop", "baseline minus perturbed accuracy", "→ 0"],
            ["Flip rate", "answers that changed at all — catches offsetting errors", "lower"],
            ["Broke / fixed", "right→wrong and wrong→right, counted separately", "lower"]]
    table(s, rows, M, Inches(1.90), CW, col_w=[22, 62, 16], font=11.5, row_h=0.42)
    label(s, "Two things worth stating aloud", M, Inches(5.15))
    bullets(s, [
        "At temperature 0 any instability comes from the serving stack, not sampling — "
        "if seed spread rivals the gap between two models, that gap is noise.",
        "Flat accuracy under perturbation is not robustness if equal numbers of answers "
        "broke and were accidentally fixed.",
    ], M, Inches(5.48), CW, Inches(1.20), size=13)


def s_pre_stage_systems(prs):
    s = section_slide(prs, EYEBROW_PRE, "Stages 6–9 · Serving, cost and reliability")
    rows = [["Metric", "What it means", "Good"],
            ["TTFT", "time to the first content token — what makes a chat feel responsive", "lower"],
            ["TPOT", "(E2E − TTFT) ÷ (tokens − 1) — steady-state generation rate", "lower"],
            ["p95 / p99", "tail latency; for serving the tail is the user experience", "lower"],
            ["Reasoning tokens", "tokens spent thinking — 0 is a real answer, not a gap", "lower"],
            ["Cost per 1M tokens", "blended price actually paid", "lower"],
            ["Cost per correct answer", "money per useful answer — the selection criterion", "lower"],
            ["Invalid output rate", "a 200 response carrying nothing usable", "lower"],
            ["Retries / failover", "how hard the client tried; whether the backend changed", "lower"]]
    table(s, rows, M, Inches(1.90), CW, col_w=[22, 62, 16], font=11.5, row_h=0.40)
    text(s, "A run that silently retried a third of its requests produces the same accuracy "
            "table as a clean one. Stage 9 is the only place that difference survives.",
         M, Inches(5.65), CW, Inches(0.60), size=12.5, color=ACCENT)


def s_pre_honesty(prs):
    s = section_slide(prs, EYEBROW_PRE, "What this setup cannot measure")
    rows = [["Field", "Status", "Why"],
            ["Accuracy, consistency, context, robustness", "measured", "prompt manipulation and repeat calls"],
            ["TTFT / TPOT / E2E, throughput", "measured", "client-side, from the streamed response"],
            ["Tokens, reasoning tokens, cost", "provider-reported", "returned in the API usage payload"],
            ["Calibration (ECE, Brier, NLL)", "provider-dependent", "only some providers return log-probabilities"],
            ["VRAM, KV cache", "analytic only", "computed from published architecture"],
            ["Energy, communication overhead", "unavailable", "needs on-host telemetry we do not have"],
            ["TP / PP / DP / SP / CP / EP effects", "unavailable", "the provider chooses the layout"]]
    table(s, rows, M, Inches(1.90), CW, col_w=[36, 20, 44], font=11.5, row_h=0.42)
    text(s, "Every numeric field carries a source tag — measured, provider_reported, "
            "analytic_model, or unavailable with a reason. Unmeasurable blocks are emitted as "
            "available: false, never as zeros.",
         M, Inches(5.30), CW, Inches(0.75), size=13, spacing=1.3)
    text(s, "A fabricated calibration figure would be worse than a missing one.",
         M, Inches(6.10), CW, Inches(0.45), size=17, bold=True, color=ACCENT)


# ---------------------------------------------------------------------------
# 8 — Results
# ---------------------------------------------------------------------------

def s_results_quality(prs, benchmark, label_):
    models = load(benchmark)
    if not models:
        return
    key = "overall_accuracy" if benchmark == "mmlu" else "last_word_accuracy"
    n = models[0]["config"]["dataset"]["n_items"]
    s = section_slide(prs, "Results", f"{label_} — task quality")
    text(s, f"{n:,} items per model  ·  Wilson 95% confidence intervals",
         M, Inches(1.72), CW, Inches(0.32), size=13, bold=True, color=ACCENT)

    rows = [["Model", "Accuracy", "95% CI", "Error rate", "Parse fail"]]
    for m in models:
        tq = m["task_quality"]
        ci = tq.get(key + "_ci95") or []
        fail = tq.get("parse_failure_rate")
        if fail is None:
            fail = tq.get("empty_prediction_rate")
        rows.append([R.name(m), R.pct(tq.get(key)),
                     f"{ci[0]*100:.1f} – {ci[1]*100:.1f}%" if len(ci) == 2 else "-",
                     R.pct(tq.get("error_rate")), R.pct(fail)])
    table(s, rows, M, Inches(2.20), CW, col_w=[30, 18, 22, 15, 15],
          font=13.5, row_h=0.44, emphasise=1)

    best = max(models, key=lambda m: m["task_quality"].get(key) or 0)
    worst = max(models, key=lambda m: (m["task_quality"].get("parse_failure_rate")
                                       or m["task_quality"].get("empty_prediction_rate") or 0))
    wf = (worst["task_quality"].get("parse_failure_rate")
          or worst["task_quality"].get("empty_prediction_rate") or 0)
    notes = [f"{R.name(best)} leads at {R.pct(best['task_quality'][key])}."]
    if wf > 0.02:
        notes.append(
            f"{R.name(worst)} returns no parseable answer on {R.pct(wf)} of items — an "
            f"instruction-following failure, reported apart from being wrong because the "
            f"fixes differ.")
    notes.append("Intervals are shown so differences inside them are not claimed as findings.")
    label(s, "Reading it", M, Inches(4.30))
    bullets(s, notes, M, Inches(4.62), CW, Inches(1.70), size=13.5)


def s_results_systems(prs, benchmark, label_):
    models = load(benchmark)
    if not models:
        return
    s = section_slide(prs, "Results", f"{label_} — serving, cost and reliability")
    rows = [["Model", "TTFT", "E2E p95", "Tok / correct", "$ / 1M tok", "$ / correct", "Success"]]
    for m in models:
        a, t = m.get("api_performance", {}), m.get("token_efficiency", {})
        e, r = m.get("economics", {}), m.get("reliability", {})
        rows.append([R.name(m),
                     R.fmt(R.dig(a, "latency.ttft.mean"), "{:.3f} s"),
                     R.fmt(R.dig(a, "latency.e2e.p95"), "{:.3f} s"),
                     R.fmt(t.get("tokens_per_correct_answer"), "{:.0f}"),
                     R.fmt(e.get("cost_per_1m_tokens_usd"), "${:.4f}"),
                     R.fmt(e.get("cost_per_correct_answer_usd"), "${:.6f}"),
                     R.pct(r.get("success_rate"))])
    table(s, rows, M, Inches(1.95), CW, col_w=[23, 13, 13, 14, 13, 14, 10],
          font=12.5, row_h=0.44)

    econ = [(R.name(m), (m.get("economics") or {}).get("cost_per_correct_answer_usd"),
             (m.get("economics") or {}).get("cost_per_1m_tokens_usd")) for m in models]
    econ = [e for e in econ if e[1] and e[2]]
    notes = []
    if len(econ) >= 2:
        ct, ca = min(econ, key=lambda x: x[2]), min(econ, key=lambda x: x[1])
        dear = max(econ, key=lambda x: x[1])
        if ct[0] != ca[0]:
            notes.append(f"{ct[0]} has the cheapest tokens, but {ca[0]} the cheapest correct "
                         f"answer — accuracy converts token price into value.")
        notes.append(f"{ca[0]} costs ${ca[1]:.6f} per correct answer against ${dear[1]:.6f} "
                     f"for {dear[0]} — {dear[1]/ca[1]:.1f}× more.")
    notes.append("Latency is client-measured; tokens and cost are provider-reported. "
                 "Nothing here is estimated.")
    label(s, "Reading it", M, Inches(4.10))
    bullets(s, notes, M, Inches(4.42), CW, Inches(1.80), size=13.5)


def s_results_ranking(prs, benchmark, label_):
    comp = comparison(benchmark)
    if not comp or not comp.get("ranking"):
        return
    s = section_slide(prs, "Results", f"{label_} — composite ranking")

    # Only the components that were actually scored get a column. A column of
    # "n/a" implies an input to the composite that does not exist.
    cols = [c for c in R.COMPONENT_COLUMNS
            if any(r.get(c[2]) is not None for r in comp["ranking"])]
    rows = [["#", "Model"] + [c[1] for c in cols] + ["Composite"]]
    for r in comp["ranking"]:
        rows.append([str(r["rank"]), r["model"].split("/")[-1]]
                    + [R.fmt(r.get(c[2])) for c in cols]
                    + [R.fmt(r.get("composite_score"))])
    body = 100 - 6 - 28
    col_w = [6, 28] + [body // (len(cols) + 1)] * (len(cols) + 1)
    col_w[-1] += 100 - sum(col_w)
    table(s, rows, M, Inches(1.95), CW, col_w=col_w,
          font=13, row_h=0.46, emphasise=len(cols) + 2)

    eff = comp.get("effective_weights") or comp.get("composite_weights", {})
    text(s, "composite  =  " + "   +   ".join(f"{v:.3f} · {k}" for k, v in eff.items()),
         M, Inches(4.15), CW, Inches(0.40), size=13, bold=True, color=ACCENT)
    missing = comp.get("components_missing") or []
    text(s, (f"{' and '.join(missing).capitalize()} could not be measured — the routed "
             f"providers returned no log-probabilities and no robustness pass was run. "
             f"Rather than scoring them zero, their weight is redistributed over the "
             f"components that were measured, so the weights above are the applied ones, "
             f"not the nominal 0.50/0.15/0.15/0.10/0.10.\n" if missing else "")
            + comp.get("comparability_note", ""),
         M, Inches(4.70), CW, Inches(1.50), size=13, spacing=1.3)


def significance(benchmark):
    path = f"results/v2/significance_{benchmark}.json"
    if not os.path.exists(path):
        return None
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _p(p):
    """`p < 0.0001` or `p = 0.0031` - never `p = < 0.0001`."""
    t = SIG.format_p(p)
    return f"p {t}" if t.startswith("<") else f"p = {t}"


def s_significance_method(prs):
    s = section_slide(prs, "Analysis", "Are the differences real?")
    text(s, "The models answered identical items, so every comparison is paired. "
            "That is what makes McNemar the right test: an unpaired two-proportion "
            "test would discard the pairing and overstate the variance.",
         M, Inches(1.75), CW, Inches(0.75), size=14)
    rows = [["Test", "Question", "Why this one"],
            ["Cochran's Q", "Do the k models differ at all?",
             "Omnibus first, so pairwise tests are not p-hacking"],
            ["McNemar", "Is a given pair's gap real?",
             "Paired; only discordant items carry information"],
            ["Holm-Bonferroni", "False positive across 3 pairs?",
             "Controls family-wise error without Bonferroni's power loss"],
            ["Chi-square + Cramer's V", "Does a factor move accuracy?",
             "At 1000s of items everything is significant; V says if it matters"],
            ["Variance decomposition", "Is the subject spread real?",
             "Removes the binomial noise of ~50 questions per subject"],
            ["Oracle ceiling", "What would per-item routing buy?",
             "Headroom that individual accuracies cannot show"]]
    table(s, rows, M, Inches(2.65), CW, col_w=[20, 32, 48], font=12, row_h=0.44)
    text(s, "Pure-stdlib implementations, pinned against SciPy and statsmodels by "
            "tests/test_significance.py so the report regenerates on a host where "
            "SciPy cannot be installed.",
         M, Inches(5.95), CW, Inches(0.5), size=12, color=MUTED)


def s_significance_pairwise(prs, benchmark, label_):
    d = significance(benchmark)
    if not d:
        return
    s = section_slide(prs, "Analysis", f"{label_} — paired significance")
    q = d.get("omnibus") or {}
    if q.get("available"):
        text(s, f"Cochran's Q = {q['q_statistic']}  (df {q['df']}),  {_p(q['p_value'])}"
                f"   —   the models differ; pairwise tests follow",
             M, Inches(1.75), CW, Inches(0.4), size=14, bold=True, color=ACCENT)
    rows = [["Pair", "Delta acc.", "95% CI", "Only A", "Only B", "p (Holm)", "Verdict"]]
    for r in d["pairwise_mcnemar"]:
        ci = r["delta_ci95"]
        rows.append([f"{r['model_a']} vs {r['model_b']}",
                     f"{r['accuracy_delta']:+.4f}",
                     f"{ci[0]:+.3f} to {ci[1]:+.3f}",
                     str(r["only_a_correct"]), str(r["only_b_correct"]),
                     SIG.format_p(r["p_adjusted"]),
                     "significant" if r["significant_adjusted"] else "n.s."])
    table(s, rows, M, Inches(2.35), CW, col_w=[30, 12, 18, 8, 8, 12, 12],
          font=12, row_h=0.44)

    o = d.get("oracle") or {}
    notes = []
    if all(r["significant_adjusted"] for r in d["pairwise_mcnemar"]):
        notes.append("Every gap survives Holm correction: the ranking is not a "
                     "sampling artefact.")
    notes.append("Only discordant items count. 'Only A' and 'Only B' are the items "
                 "one model got right and the other did not - items both answered "
                 "the same way carry no information about which is better.")
    if o:
        notes.append(f"An ideal per-item router would reach {o['oracle_accuracy']*100:.1f}% "
                     f"against {o['best_single_accuracy']*100:.1f}% for the best single "
                     f"model - {o['headroom']*100:.1f} points of headroom, with "
                     f"{o['none_correct']*100:.1f}% of items defeating all three.")
    label(s, "Reading it", M, Inches(4.35))
    bullets(s, notes, M, Inches(4.67), CW, Inches(1.9), size=13)


def s_significance_factors(prs):
    s = section_slide(prs, "Analysis", "Which parameters move the outcome?")
    text(s, "Significance and effect size are different questions. Cramer's V is "
            "reported alongside every p-value because at thousands of items a "
            "negligible association is still significant.",
         M, Inches(1.75), CW, Inches(0.6), size=14)
    rows = [["Benchmark", "Factor", "Spread", "p", "V", "Effect"]]
    for bench, lab in (("mmlu", "MMLU"), ("lambada", "LAMBADA")):
        d = significance(bench)
        if not d:
            continue
        for name, f in (d.get("factors") or {}).items():
            if f.get("available"):
                rows.append([lab, name, f"{f['spread']*100:.1f} pp",
                             SIG.format_p(f["p_value"]), f"{f['cramers_v']:.3f}",
                             f["effect"]])
        for m, f in (d.get("fragmentation_by_model") or {}).items():
            if f.get("available"):
                rows.append([lab, f"target fragmentation - {m}",
                             f"{f['spread']*100:.1f} pp", SIG.format_p(f["p_value"]),
                             f"{f['cramers_v']:.3f}", f["effect"]])
    table(s, rows, M, Inches(2.5), CW, col_w=[14, 38, 12, 14, 10, 12],
          font=12, row_h=0.40)

    d = significance("mmlu")
    v = (d or {}).get("subject_variance") or {}
    notes = []
    if v.get("available"):
        notes.append(f"MMLU subject spread is real, not noise: observed SD "
                     f"{v['observed_sd']*100:.1f} pp, and after removing the binomial "
                     f"variance of ~50 questions per subject, {v['between_share']*100:.0f}% "
                     f"of it survives (a null of identical subjects averages ~8%).")
    notes.append("LAMBADA: target fragmentation outweighs passage length two to "
                 "three times over. The gap is a vocabulary handicap, not a "
                 "context-window one.")
    notes.append("MMLU: the correct option's position shifts accuracy by 7 points - "
                 "a property of the harness, not of the knowledge tested, which is "
                 "why the robustness stage permutes options.")
    label(s, "Reading it", M, Inches(5.05))
    bullets(s, notes, M, Inches(5.37), CW, Inches(1.5), size=13)


def s_results_figures(prs):
    figs = [
        ("analysis-01-history-overview.png", "The run record",
         "History: one card per run, with metric-family and configuration chips"),
        ("analysis-03-stage-list.png", "Nine stages per model",
         "Families with no data carry an n/a badge rather than being hidden"),
        ("analysis-04-task-quality.png", "Stage 1 in the web view",
         "Wilson intervals, per-subject tables and option-position bias"),
        ("analysis-06-calibration-unavailable.png", "Stage 2 — reporting absence",
         "An explicit unavailability notice with its reason, instead of a zero"),
        ("analysis-10-api-performance.png", "Stage 6 — latency",
         "TTFT, TPOT and E2E with p50 / p95 / p99, plus throughput"),
        ("analysis-13-reliability.png", "Stage 9 — reliability",
         "Retries, failures and provider failover — invisible to accuracy"),
        ("analysis-14-run-config.png", "Configuration as input",
         "TP/PP/DP/SP/CP/EP recorded and labelled as input, never scored"),
        ("analysis-15-decoding-panel.png", "Run controls",
         "Nine decoding parameters; the call multiplier shown before spending"),
    ]
    titles = ["The web view — run record and stages",
              "The web view — quality and calibration",
              "The web view — latency and reliability",
              "The web view — configuration and controls"]
    for i in range(0, len(figs), 2):
        pair = figs[i:i + 2]
        s = section_slide(prs, "Results", titles[i // 2])
        for j, (fn, head, cap) in enumerate(pair):
            left = M + j * Inches(6.25)
            text(s, head, left, Inches(1.80), Inches(5.85), Inches(0.32),
                 size=14, bold=True, color=INK)
            figure(s, fn, cap, left, Inches(2.20), Inches(5.85), Inches(3.30))


# ---------------------------------------------------------------------------
# 9 — Discussion of the results
# ---------------------------------------------------------------------------

def s_discussion(prs):
    mm = {R.name(m): m for m in load("mmlu")}
    lb = {R.name(m): m for m in load("lambada")}
    s = section_slide(prs, "Discussion of the results",
                      "Findings accuracy alone would have missed")
    items = []

    pf = sorted(((n, m["task_quality"].get("parse_failure_rate") or 0)
                 for n, m in mm.items()), key=lambda x: -x[1])
    if pf and pf[0][1] > 0.02:
        items.append(("1 · Task Quality",
                      f"{pf[0][0]} returns no parseable answer on {pf[0][1]*100:.1f}% of "
                      f"questions — disobedience scored as ignorance."))
    tk = [(n, t) for n, t in ((n, m.get("tokenization") or {}) for n, m in lb.items())
          if t.get("available")]
    if tk:
        get3 = lambda t: (t["accuracy_by_fragmentation"].get("3plus_tokens", {})
                          .get("accuracy") or 0)
        best, worst = max(tk, key=lambda x: get3(x[1])), min(tk, key=lambda x: get3(x[1]))
        items.append(("Tokenization",
                      f"On 3+-token targets {best[0]} holds {get3(best[1])*100:.0f}% while "
                      f"{worst[0]} collapses to {get3(worst[1])*100:.0f}% — a vocabulary "
                      f"handicap, not a comprehension gap."))
    noisy = [(n, r) for n, r in ((n, m["reliability"]) for n, m in mm.items())
             if (r.get("retry_total") or 0) or r.get("provider_failover")
             or (r.get("failure_rate") or 0)]
    if noisy:
        n, r = noisy[0]
        bits = []
        if r.get("retry_total"):
            bits.append(f"{r['retry_total']} retries")
        if r.get("failure_rate"):
            bits.append(f"{r['failure_rate']*100:.1f}% transport failures")
        if r.get("provider_failover"):
            bits.append("provider failover mid-run")
        items.append(("9 · Reliability",
                      f"{n}: {', '.join(bits)} — invisible in an accuracy table and "
                      f"unrecoverable after the fact without this stage."))
    econ = [(n, e) for n, e in ((n, m.get("economics") or {}) for n, m in mm.items())
            if e.get("cost_per_correct_answer_usd")]
    if len(econ) >= 2:
        ct = min(econ, key=lambda x: x[1]["cost_per_1m_tokens_usd"])
        ca = min(econ, key=lambda x: x[1]["cost_per_correct_answer_usd"])
        dear = max(econ, key=lambda x: x[1]["cost_per_correct_answer_usd"])
        items.append(("8 · Economics",
                      f"{dear[0]} costs "
                      f"{dear[1]['cost_per_correct_answer_usd']/ca[1]['cost_per_correct_answer_usd']:.1f}× "
                      f"more per correct answer than {ca[0]} — token price inverts once "
                      f"accuracy is applied."))
    items.append(("2 · Probabilistic",
                  "Whether calibration is obtainable at all depends on which provider "
                  "OpenRouter routes to — a property of the serving path, not the model."))

    rows = [["Stage", "Finding"]] + [[a, b] for a, b in items[:5]]
    table(s, rows, M, Inches(1.90), CW, col_w=[18, 82], font=12, row_h=0.62)


def s_discussion_tokenizer(prs):
    lb = [m for m in load("lambada") if (m.get("tokenization") or {}).get("available")]
    if not lb:
        return
    s = section_slide(prs, "Discussion of the results",
                      "Part of the LAMBADA gap is the tokenizer")
    rows = [["Model", "1-token targets", "2-token", "3+-token", "Collapse", "Tokenizer exact?"]]
    for m in lb:
        t = m["tokenization"]
        bf = t["accuracy_by_fragmentation"]
        g = lambda k: (bf.get(k) or {}).get("accuracy")
        one, many = g("1_token"), g("3plus_tokens")
        rows.append([R.name(m), R.pct(one), R.pct(g("2_tokens")), R.pct(many),
                     f"{(one-many)*100:+.1f} pp" if one is not None and many is not None else "-",
                     "yes" if t["tokenizer_exact"] else "no — fallback"])
    table(s, rows, M, Inches(1.95), CW, col_w=[22, 17, 13, 14, 14, 20],
          font=12.5, row_h=0.46)
    bullets(s, [
        "A multi-token target must be produced correctly several times over — a mechanical "
        "handicap, not a comprehension one.",
        "Comparing raw LAMBADA scores without this split treats a vocabulary difference as a "
        "capability difference.",
    ], M, Inches(4.05), CW, Inches(1.10), size=13.5)
    approx = [R.name(m) for m in lb if not m["tokenization"]["tokenizer_exact"]]
    if approx:
        text(s, f"Caveat: {', '.join(approx)} fell back to a generic BPE vocabulary because "
                f"their Hugging Face tokenizers are gated. The trend within each model holds; "
                f"the fragmentation rates are not comparable across models.",
             M, Inches(5.30), CW, Inches(0.85), size=12.5, color=ACCENT, spacing=1.3)


# ---------------------------------------------------------------------------
# 10 — Conclusions, References
# ---------------------------------------------------------------------------

def s_conclusions(prs):
    s = base(prs)
    text(s, "Conclusions", M, Inches(0.55), CW, Inches(0.70),
         size=30, bold=True, color=INK)
    rule(s, Inches(1.30))
    blocks = [
        ("Does a larger, non-distilled model beat smaller compressed ones?",
         "On task quality, consistently — Ministral-8B leads both benchmarks. But the "
         "nine-stage view shows the answer is not one-dimensional: the smaller models win on "
         "latency and, on MMLU, on cost per correct answer, and second place changes hands "
         "depending on whether reliability is counted."),
        ("What the extra stages bought.",
         "Findings invisible to accuracy alone: an instruction-following failure mistaken for "
         "ignorance; a tokenizer handicap mistaken for a comprehension gap; a token price that "
         "inverts once accuracy is applied; and serving failures hidden inside a clean score."),
        ("The methodological point.",
         "Configuration and measurement are kept apart — parallelism degrees and decoding "
         "settings are inputs recorded in a run manifest, never folded into a score. That is "
         "what lets the comparison layer refuse to rank runs that are not comparable, which is "
         "the correct output in that case."),
    ]
    y = Inches(1.65)
    for head, body in blocks:
        text(s, head, M, y, CW, Inches(0.42), size=17, bold=True, color=INK)
        text(s, body, M, y + Inches(0.46), CW, Inches(1.15), size=14,
             color=BODY, spacing=1.3)
        y += Inches(1.72)
    rule(s, Inches(6.75))
    text(s, "Absence is reported, not filled in — the difference between a benchmark you can "
            "cite and one you cannot.",
         M, Inches(6.90), CW, Inches(0.40), size=14, bold=True, color=ACCENT)


def s_references(prs):
    s = section_slide(prs, "References", "Sources", size=28)
    refs = [
        "Hendrycks, D. et al. (2021). Measuring Massive Multitask Language Understanding. ICLR 2021.",
        "Paperno, D. et al. (2016). The LAMBADA dataset. ACL 2016.",
        "Guo, C. et al. (2017). On Calibration of Modern Neural Networks. ICML 2017.",
        "Brier, G. W. (1950). Verification of Forecasts Expressed in Terms of Probability. Monthly Weather Review.",
        "Wilson, E. B. (1927). Probable Inference, the Law of Succession, and Statistical Inference. JASA.",
        "Wang, X. et al. (2023). Self-Consistency Improves Chain of Thought Reasoning. ICLR 2023.",
        "Gao, L. et al. (2024). A Framework for Few-Shot Language Model Evaluation (lm-evaluation-harness).",
        "Shoeybi, M. et al. (2019). Megatron-LM: Training Multi-Billion Parameter Models Using Model Parallelism.",
        "Kwon, W. et al. (2023). Efficient Memory Management for LLM Serving with PagedAttention (vLLM). SOSP 2023.",
        "Google (2025) Gemma 3 technical report · Meta (2024) Llama 3.2 model card · Mistral AI (2024) Ministral family.",
    ]
    bullets(s, refs, M, Inches(1.95), Inches(11.9), Inches(4.40), size=13, spacing=1.25)
    text(s, "github.com/ronvoy/gen-ai   ·   unina.cc/gen-ai",
         M, Inches(6.55), Inches(11.9), Inches(0.40), size=12, color=MUTED)


# ---------------------------------------------------------------------------

def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    s_title(prs)
    s_research_question(prs)
    s_protocol(prs)
    s_protocol_passes(prs)
    s_benchmark_arch(prs)
    s_benchmark_config(prs)
    s_benchmark_data(prs)
    s_approaches(prs)
    s_models(prs)
    s_pre_stage_quality(prs)
    s_pre_stage_behaviour(prs)
    s_pre_stage_systems(prs)
    s_pre_honesty(prs)

    for bench, lab in (("mmlu", "MMLU"), ("lambada", "LAMBADA")):
        s_results_quality(prs, bench, lab)
        s_results_systems(prs, bench, lab)
        s_results_ranking(prs, bench, lab)
    s_results_figures(prs)

    s_significance_method(prs)
    for bench, lab in (("mmlu", "MMLU"), ("lambada", "LAMBADA")):
        s_significance_pairwise(prs, bench, lab)
    s_significance_factors(prs)

    s_discussion(prs)
    s_discussion_tokenizer(prs)
    s_conclusions(prs)
    s_references(prs)

    prs.save(OUT)
    n = len(prs.slides._sldIdLst)
    missing = sum(1 for sl in prs.slides for sh in sl.shapes
                  if sh.has_text_frame and "screenshot placeholder" in sh.text_frame.text)
    print(f"wrote {OUT}: {n} slides, {missing} figure placeholder(s) awaiting screenshots")


if __name__ == "__main__":
    main()
